"""
Requirements Analyzer Agent

This agent analyzes requirements from multiple document formats including
DOCX, Excel, PDF, and text files. It extracts, categorizes, and visualizes
requirements for technical architecture design.
"""

from typing import List, Dict, Optional, Set, Union
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
import re
import json


class RequirementType(Enum):
    """Types of requirements"""
    FUNCTIONAL = "functional"
    NON_FUNCTIONAL = "non_functional"
    BUSINESS = "business"
    TECHNICAL = "technical"
    SECURITY = "security"
    PERFORMANCE = "performance"
    USABILITY = "usability"
    COMPLIANCE = "compliance"
    CONSTRAINT = "constraint"


class RequirementPriority(Enum):
    """Requirement priority levels"""
    CRITICAL = "critical"
    HIGH = "high"
    MEDIUM = "medium"
    LOW = "low"


class RequirementStatus(Enum):
    """Requirement status"""
    DRAFT = "draft"
    PROPOSED = "proposed"
    APPROVED = "approved"
    IMPLEMENTED = "implemented"
    VERIFIED = "verified"
    REJECTED = "rejected"


@dataclass
class Requirement:
    """Represents a single requirement"""
    id: str
    title: str
    description: str
    type: RequirementType
    priority: RequirementPriority
    status: RequirementStatus = RequirementStatus.DRAFT
    source: Optional[str] = None
    stakeholder: Optional[str] = None
    acceptance_criteria: List[str] = field(default_factory=list)
    dependencies: List[str] = field(default_factory=list)
    tags: List[str] = field(default_factory=list)
    rationale: Optional[str] = None
    estimated_effort: Optional[str] = None


@dataclass
class RequirementCategory:
    """Categorized requirements"""
    category: str
    requirements: List[Requirement]
    count: int = 0

    def __post_init__(self):
        self.count = len(self.requirements)


@dataclass
class RequirementAnalysis:
    """Complete requirement analysis results"""
    requirements: List[Requirement]
    categories: Dict[str, RequirementCategory]
    total_count: int
    summary: Dict[str, int]
    sources: List[str]
    stakeholders: Set[str]
    recommendations: List[str]


class RequirementsAnalyzer:
    """
    Requirements Analyzer Agent

    Features:
    - Parse multiple document formats (DOCX, Excel, PDF, TXT, MD)
    - Extract and categorize requirements
    - Generate requirement diagrams (use case, traceability matrix)
    - Create requirement specifications
    - Priority analysis and recommendations
    - Generate requirement documentation
    """

    def __init__(self):
        self.requirements: List[Requirement] = []
        self.requirement_patterns = {
            'functional': r'\b(shall|must|will|should)\s+(?:be able to|allow|enable|provide|support)',
            'non_functional': r'\b(performance|scalability|security|availability|reliability|maintainability)',
            'must_have': r'\b(must|shall|required|mandatory)\b',
            'should_have': r'\b(should|recommended|preferred)\b',
            'could_have': r'\b(could|optional|nice to have|may)\b',
        }

    def analyze_documents(
        self,
        file_paths: List[str],
        auto_categorize: bool = True
    ) -> RequirementAnalysis:
        """
        Analyze multiple requirement documents.

        Args:
            file_paths: List of document paths
            auto_categorize: Automatically categorize requirements

        Returns:
            RequirementAnalysis object
        """
        all_requirements = []
        sources = []

        for file_path in file_paths:
            path = Path(file_path)
            sources.append(path.name)

            if not path.exists():
                print(f"Warning: File not found: {file_path}")
                continue

            # Parse based on file extension
            if path.suffix.lower() == '.docx':
                reqs = self._parse_docx(file_path)
            elif path.suffix.lower() in ['.xlsx', '.xls']:
                reqs = self._parse_excel(file_path)
            elif path.suffix.lower() == '.pdf':
                reqs = self._parse_pdf(file_path)
            elif path.suffix.lower() in ['.txt', '.md']:
                reqs = self._parse_text(file_path)
            elif path.suffix.lower() == '.json':
                reqs = self._parse_json(file_path)
            else:
                print(f"Warning: Unsupported file type: {path.suffix}")
                continue

            # Add source to each requirement
            for req in reqs:
                req.source = path.name

            all_requirements.extend(reqs)

        # Auto-categorize if enabled
        if auto_categorize:
            all_requirements = self._auto_categorize(all_requirements)

        # Create analysis
        analysis = self._create_analysis(all_requirements, sources)

        return analysis

    def _parse_docx(self, file_path: str) -> List[Requirement]:
        """Parse DOCX file (requires python-docx)"""
        requirements = []

        try:
            from docx import Document
            doc = Document(file_path)

            req_id = 1
            current_section = None

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Check if it's a heading (potential section)
                if para.style.name.startswith('Heading'):
                    current_section = text
                    continue

                # Look for requirement patterns
                if self._is_requirement(text):
                    req = self._extract_requirement(
                        f"REQ-{req_id:03d}",
                        text,
                        current_section
                    )
                    requirements.append(req)
                    req_id += 1

        except ImportError:
            print("Warning: python-docx not installed. Install with: pip install python-docx")
            # Fallback: treat as text
            return self._parse_text(file_path)
        except Exception as e:
            print(f"Error parsing DOCX: {e}")

        return requirements

    def _parse_excel(self, file_path: str) -> List[Requirement]:
        """Parse Excel file (requires openpyxl)"""
        requirements = []

        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path)
            ws = wb.active

            # Assume first row is header
            headers = [cell.value for cell in ws[1]]

            # Find relevant columns
            id_col = self._find_column(headers, ['id', 'req id', 'requirement id'])
            title_col = self._find_column(headers, ['title', 'name', 'requirement'])
            desc_col = self._find_column(headers, ['description', 'details', 'desc'])
            type_col = self._find_column(headers, ['type', 'category'])
            priority_col = self._find_column(headers, ['priority', 'importance'])
            status_col = self._find_column(headers, ['status', 'state'])

            # Parse rows
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue

                req_id = str(row[id_col]) if id_col is not None and row[id_col] else f"REQ-{len(requirements)+1:03d}"
                title = str(row[title_col]) if title_col is not None and row[title_col] else ""
                description = str(row[desc_col]) if desc_col is not None and row[desc_col] else title

                # Determine type
                req_type = RequirementType.FUNCTIONAL
                if type_col is not None and row[type_col]:
                    req_type = self._parse_requirement_type(str(row[type_col]))

                # Determine priority
                priority = RequirementPriority.MEDIUM
                if priority_col is not None and row[priority_col]:
                    priority = self._parse_priority(str(row[priority_col]))

                # Determine status
                status = RequirementStatus.DRAFT
                if status_col is not None and row[status_col]:
                    status = self._parse_status(str(row[status_col]))

                req = Requirement(
                    id=req_id,
                    title=title,
                    description=description,
                    type=req_type,
                    priority=priority,
                    status=status
                )
                requirements.append(req)

        except ImportError:
            print("Warning: openpyxl not installed. Install with: pip install openpyxl")
        except Exception as e:
            print(f"Error parsing Excel: {e}")

        return requirements

    def _parse_pdf(self, file_path: str) -> List[Requirement]:
        """Parse PDF file (requires PyPDF2 or pdfplumber)"""
        requirements = []
        text_content = ""

        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text_content += page.extract_text() + "\n"
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(file_path)
                for page in reader.pages:
                    text_content += page.extract_text() + "\n"
            except ImportError:
                print("Warning: Neither pdfplumber nor PyPDF2 installed.")
                print("Install with: pip install pdfplumber")
                return requirements
        except Exception as e:
            print(f"Error parsing PDF: {e}")
            return requirements

        # Process extracted text
        lines = text_content.split('\n')
        req_id = 1
        current_section = None

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Check for section headers
            if line.isupper() or line.endswith(':'):
                current_section = line.rstrip(':')
                continue

            # Look for requirements
            if self._is_requirement(line):
                req = self._extract_requirement(
                    f"REQ-{req_id:03d}",
                    line,
                    current_section
                )
                requirements.append(req)
                req_id += 1

        return requirements

    def _parse_text(self, file_path: str) -> List[Requirement]:
        """Parse text or markdown file"""
        requirements = []

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            lines = content.split('\n')
            req_id = 1
            current_section = None

            for line in lines:
                line = line.strip()
                if not line:
                    continue

                # Check for markdown headers
                if line.startswith('#'):
                    current_section = line.lstrip('#').strip()
                    continue

                # Look for numbered requirements
                if re.match(r'^\d+\.', line):
                    line = re.sub(r'^\d+\.\s*', '', line)

                # Look for requirements
                if self._is_requirement(line):
                    req = self._extract_requirement(
                        f"REQ-{req_id:03d}",
                        line,
                        current_section
                    )
                    requirements.append(req)
                    req_id += 1

        except Exception as e:
            print(f"Error parsing text file: {e}")

        return requirements

    def _parse_json(self, file_path: str) -> List[Requirement]:
        """Parse JSON requirements file"""
        requirements = []

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)

            if isinstance(data, list):
                for item in data:
                    req = self._requirement_from_dict(item)
                    requirements.append(req)
            elif isinstance(data, dict) and 'requirements' in data:
                for item in data['requirements']:
                    req = self._requirement_from_dict(item)
                    requirements.append(req)

        except Exception as e:
            print(f"Error parsing JSON: {e}")

        return requirements

    def _is_requirement(self, text: str) -> bool:
        """Check if text contains a requirement"""
        # Look for requirement keywords
        keywords = ['shall', 'must', 'will', 'should', 'require', 'need']
        text_lower = text.lower()

        return any(keyword in text_lower for keyword in keywords) and len(text) > 20

    def _extract_requirement(
        self,
        req_id: str,
        text: str,
        section: Optional[str] = None
    ) -> Requirement:
        """Extract requirement from text"""
        # Determine type based on text analysis
        req_type = self._analyze_requirement_type(text)

        # Determine priority based on keywords
        priority = self._analyze_priority(text)

        # Extract title (first sentence or first 100 chars)
        title = text.split('.')[0][:100] if '.' in text else text[:100]

        return Requirement(
            id=req_id,
            title=title,
            description=text,
            type=req_type,
            priority=priority,
            status=RequirementStatus.DRAFT,
            tags=[section] if section else []
        )

    def _analyze_requirement_type(self, text: str) -> RequirementType:
        """Analyze requirement type from text"""
        text_lower = text.lower()

        # Security keywords
        if any(kw in text_lower for kw in ['security', 'authentication', 'authorization', 'encryption', 'secure']):
            return RequirementType.SECURITY

        # Performance keywords
        if any(kw in text_lower for kw in ['performance', 'speed', 'response time', 'latency', 'throughput']):
            return RequirementType.PERFORMANCE

        # Usability keywords
        if any(kw in text_lower for kw in ['usability', 'user interface', 'ui', 'ux', 'user experience']):
            return RequirementType.USABILITY

        # Compliance keywords
        if any(kw in text_lower for kw in ['compliance', 'regulation', 'gdpr', 'hipaa', 'legal']):
            return RequirementType.COMPLIANCE

        # Technical keywords
        if any(kw in text_lower for kw in ['api', 'database', 'integration', 'system', 'technical']):
            return RequirementType.TECHNICAL

        # Default to functional
        return RequirementType.FUNCTIONAL

    def _analyze_priority(self, text: str) -> RequirementPriority:
        """Analyze priority from text"""
        text_lower = text.lower()

        if any(kw in text_lower for kw in ['must', 'critical', 'essential', 'required', 'shall']):
            return RequirementPriority.HIGH
        elif any(kw in text_lower for kw in ['should', 'important', 'recommended']):
            return RequirementPriority.MEDIUM
        elif any(kw in text_lower for kw in ['could', 'optional', 'nice to have', 'may']):
            return RequirementPriority.LOW

        return RequirementPriority.MEDIUM

    def _auto_categorize(self, requirements: List[Requirement]) -> List[Requirement]:
        """Automatically categorize requirements"""
        # Already done in extraction, but can be refined here
        return requirements

    def _create_analysis(
        self,
        requirements: List[Requirement],
        sources: List[str]
    ) -> RequirementAnalysis:
        """Create requirement analysis"""
        # Categorize by type
        categories = {}
        for req_type in RequirementType:
            reqs = [r for r in requirements if r.type == req_type]
            if reqs:
                categories[req_type.value] = RequirementCategory(
                    category=req_type.value,
                    requirements=reqs,
                    count=len(reqs)
                )

        # Create summary
        summary = {
            'total': len(requirements),
            'by_type': {rt.value: len([r for r in requirements if r.type == rt]) for rt in RequirementType},
            'by_priority': {rp.value: len([r for r in requirements if r.priority == rp]) for rp in RequirementPriority},
        }

        # Extract stakeholders
        stakeholders = set(r.stakeholder for r in requirements if r.stakeholder)

        # Generate recommendations
        recommendations = self._generate_recommendations(requirements)

        return RequirementAnalysis(
            requirements=requirements,
            categories=categories,
            total_count=len(requirements),
            summary=summary,
            sources=sources,
            stakeholders=stakeholders,
            recommendations=recommendations
        )

    def _generate_recommendations(self, requirements: List[Requirement]) -> List[str]:
        """Generate recommendations based on requirements"""
        recommendations = []

        # Check for missing acceptance criteria
        missing_ac = [r for r in requirements if not r.acceptance_criteria]
        if len(missing_ac) > len(requirements) * 0.5:
            recommendations.append("Add acceptance criteria to requirements for better testability")

        # Check priority distribution
        high_priority = len([r for r in requirements if r.priority == RequirementPriority.HIGH])
        if high_priority > len(requirements) * 0.7:
            recommendations.append("Too many high-priority requirements. Consider reprioritizing")

        # Check for security requirements
        security_reqs = [r for r in requirements if r.type == RequirementType.SECURITY]
        if len(security_reqs) < len(requirements) * 0.1:
            recommendations.append("Consider adding more security requirements")

        # Check for performance requirements
        perf_reqs = [r for r in requirements if r.type == RequirementType.PERFORMANCE]
        if len(perf_reqs) < len(requirements) * 0.05:
            recommendations.append("Consider defining performance requirements")

        return recommendations

    def _find_column(self, headers: List, keywords: List[str]) -> Optional[int]:
        """Find column index by keywords"""
        for i, header in enumerate(headers):
            if header and any(kw in str(header).lower() for kw in keywords):
                return i
        return None

    def _parse_requirement_type(self, type_str: str) -> RequirementType:
        """Parse requirement type from string"""
        type_str_lower = type_str.lower()
        for req_type in RequirementType:
            if req_type.value in type_str_lower:
                return req_type
        return RequirementType.FUNCTIONAL

    def _parse_priority(self, priority_str: str) -> RequirementPriority:
        """Parse priority from string"""
        priority_str_lower = priority_str.lower()
        for priority in RequirementPriority:
            if priority.value in priority_str_lower:
                return priority
        return RequirementPriority.MEDIUM

    def _parse_status(self, status_str: str) -> RequirementStatus:
        """Parse status from string"""
        status_str_lower = status_str.lower()
        for status in RequirementStatus:
            if status.value in status_str_lower:
                return status
        return RequirementStatus.DRAFT

    def _requirement_from_dict(self, data: Dict) -> Requirement:
        """Create Requirement from dictionary"""
        return Requirement(
            id=data.get('id', 'REQ-???'),
            title=data.get('title', ''),
            description=data.get('description', ''),
            type=self._parse_requirement_type(data.get('type', 'functional')),
            priority=self._parse_priority(data.get('priority', 'medium')),
            status=self._parse_status(data.get('status', 'draft')),
            source=data.get('source'),
            stakeholder=data.get('stakeholder'),
            acceptance_criteria=data.get('acceptance_criteria', []),
            dependencies=data.get('dependencies', []),
            tags=data.get('tags', []),
            rationale=data.get('rationale'),
            estimated_effort=data.get('estimated_effort')
        )

    def generate_use_case_diagram(self, analysis: RequirementAnalysis) -> str:
        """Generate PlantUML use case diagram from requirements"""
        diagram = "@startuml\n"
        diagram += "left to right direction\n"
        diagram += "skinparam packageStyle rectangle\n\n"

        # Extract actors from stakeholders or infer from requirements
        actors = analysis.stakeholders if analysis.stakeholders else {"User", "Admin", "System"}

        # Add actors
        for actor in actors:
            diagram += f'actor "{actor}" as {actor.replace(" ", "_")}\n'

        diagram += "\n"

        # Group by category
        for category_name, category in analysis.categories.items():
            diagram += f'rectangle "{category_name.replace("_", " ").title()}" {{\n'

            for req in category.requirements[:5]:  # Limit to avoid overcrowding
                use_case_name = req.title[:50]
                use_case_id = req.id.replace("-", "_")
                diagram += f'  usecase "{use_case_name}" as {use_case_id}\n'

            diagram += "}\n\n"

        # Add relationships (actors to use cases)
        for category in analysis.categories.values():
            for req in category.requirements[:5]:
                actor = list(actors)[0].replace(" ", "_") if actors else "User"
                use_case_id = req.id.replace("-", "_")
                diagram += f'{actor} --> {use_case_id}\n'

        diagram += "\n@enduml"
        return diagram

    def generate_traceability_matrix(self, analysis: RequirementAnalysis) -> str:
        """Generate requirement traceability matrix in Markdown"""
        matrix = "# Requirements Traceability Matrix\n\n"
        matrix += "| Req ID | Title | Type | Priority | Status | Dependencies |\n"
        matrix += "|--------|-------|------|----------|--------|-------------|\n"

        for req in analysis.requirements:
            deps = ", ".join(req.dependencies) if req.dependencies else "None"
            matrix += f"| {req.id} | {req.title[:50]} | {req.type.value} | "
            matrix += f"{req.priority.value} | {req.status.value} | {deps} |\n"

        return matrix

    def generate_requirement_report(self, analysis: RequirementAnalysis) -> str:
        """Generate comprehensive requirement report in Markdown"""
        report = "# Requirements Analysis Report\n\n"

        # Summary
        report += "## Executive Summary\n\n"
        report += f"**Total Requirements**: {analysis.total_count}\n\n"
        report += f"**Sources**: {', '.join(analysis.sources)}\n\n"

        # By Type
        report += "### Requirements by Type\n\n"
        for req_type, count in analysis.summary['by_type'].items():
            if count > 0:
                report += f"- **{req_type.replace('_', ' ').title()}**: {count}\n"

        report += "\n### Requirements by Priority\n\n"
        for priority, count in analysis.summary['by_priority'].items():
            if count > 0:
                report += f"- **{priority.title()}**: {count}\n"

        # Detailed requirements
        report += "\n## Detailed Requirements\n\n"

        for category_name, category in analysis.categories.items():
            report += f"### {category_name.replace('_', ' ').title()}\n\n"

            for req in category.requirements:
                report += f"#### {req.id}: {req.title}\n\n"
                report += f"**Description**: {req.description}\n\n"
                report += f"**Priority**: {req.priority.value.title()}\n\n"
                report += f"**Status**: {req.status.value.title()}\n\n"

                if req.acceptance_criteria:
                    report += "**Acceptance Criteria**:\n"
                    for ac in req.acceptance_criteria:
                        report += f"- {ac}\n"
                    report += "\n"

                if req.dependencies:
                    report += f"**Dependencies**: {', '.join(req.dependencies)}\n\n"

                report += "---\n\n"

        # Recommendations
        if analysis.recommendations:
            report += "## Recommendations\n\n"
            for i, rec in enumerate(analysis.recommendations, 1):
                report += f"{i}. {rec}\n"

        return report


if __name__ == "__main__":
    # Example usage
    analyzer = RequirementsAnalyzer()

    # Example: Analyze requirements from multiple files
    # analysis = analyzer.analyze_documents([
    #     "requirements.docx",
    #     "requirements.xlsx",
    #     "requirements.pdf",
    #     "requirements.txt"
    # ])

    # print(f"Total requirements: {analysis.total_count}")
    # print(analyzer.generate_requirement_report(analysis))
